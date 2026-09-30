"""
06_build_excel_workbook.py

Builds stress_test/External_Stress_Test.xlsx — a fully formula-driven,
auditable reproduction of the Gate 3 stress test (04_gate3_stress_test.py),
not a static table paste. Six sheets: README, Inputs, Scenarios,
Calculations, Summary, QA_Checks. See RESEARCH_LEDGER.md D-025 for the
build record, including a row-reference bug that was found by this
workbook's own QA_Checks sheet and fixed before delivery.

No macros/VBA. No new assumptions beyond the locked Gate 3 specification.
After running, recalculate the workbook (e.g. with the xlsx skill's
recalc.py) to populate formula values and confirm QA_Checks reads
"ALL CHECKS PASSED".

Requires (must already exist — run 02_build_master_dataset.py and
04_gate3_stress_test.py first):
    data/derived/master_analytical_dataset.csv  (Inputs sheet line items)
    data/derived/gate3_stress_test.csv          (QA_Checks reference values)

The Inputs-sheet quarterly/monthly line items and the QA_Checks reference
values are read from these two files at build time rather than duplicated
as hard-coded constants in this script — if the underlying data changes and
this script is re-run, the workbook (including its blue "hardcoded" Inputs
cells and its QA reference column) regenerates in step with it. The
Excel formulas, sheet structure, and locked results are unchanged.
"""
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

BASELINE_Q = ["2025-Q2", "2025-Q3", "2025-Q4", "2026-Q1"]
BASELINE_MONTHS = ["2025-M04", "2025-M05", "2025-M06", "2025-M07", "2025-M08", "2025-M09",
                   "2025-M10", "2025-M11", "2025-M12", "2026-M01", "2026-M02", "2026-M03"]


def load_workbook_data():
    """Load the granular Inputs-sheet line items from the master dataset, and the
    baseline/result reference values from the Gate 3 stress-test output — replacing
    what used to be hard-coded Python constants duplicated from those files."""
    master = pd.read_csv("data/derived/master_analytical_dataset.csv")
    gate3 = pd.read_csv("data/derived/gate3_stress_test.csv")

    def quarterly(iso, var):
        sub = master[(master.country_iso3 == iso) & (master.variable_code == var) & (master.period.isin(BASELINE_Q))]
        sub = sub.set_index("period").reindex(BASELINE_Q)
        assert sub["value"].notna().all(), f"{iso} {var}: missing a baseline quarter"
        return sub["value"].tolist()

    quarterly_data = {}
    for iso in ["BGD", "PAK", "VNM"]:
        quarterly_data[iso] = {
            "exports": quarterly(iso, "goods_exports"),
            "imports": quarterly(iso, "goods_imports"),
            "ca": quarterly(iso, "current_account_balance"),
        }

    remittance_data = {}
    for iso in ["BGD", "PAK"]:
        rem = master[(master.country_iso3 == iso) & (master.variable_code == "remittances")].set_index("period")
        rem = rem.reindex(BASELINE_MONTHS)
        assert rem["value"].notna().all(), f"{iso} remittances: missing a baseline month"
        remittance_data[iso] = list(zip(BASELINE_MONTHS, rem["value"].tolist()))

    energy_row = master[(master.country_iso3 == "WLD") & (master.variable_code == "energy_price_index")].set_index("period")
    energy_row = energy_row.reindex(BASELINE_MONTHS)
    assert energy_row["value"].notna().all(), "energy index: missing a baseline month"
    energy_data = list(zip(BASELINE_MONTHS, energy_row["value"].tolist()))

    # QA reference values, read directly from gate3_stress_test.csv rather than retyped
    qa_baseline_refs, qa_locked_refs = {}, {}
    for iso, name in [("BGD", "Bangladesh"), ("PAK", "Pakistan"), ("VNM", "Vietnam")]:
        country_rows = gate3[gate3.country_iso3 == iso]
        first_row = country_rows.iloc[0]
        qa_baseline_refs[iso] = dict(
            exports=first_row["baseline_exports"], imports=first_row["baseline_imports"],
            ca=first_row["baseline_current_account"],
            remittances=first_row["baseline_remittances"] if iso != "VNM" else None,
        )
        combined_row = country_rows[country_rows["scenario"].str.contains("Combined")].iloc[0]
        qa_locked_refs[iso] = dict(
            deterioration=combined_row["total_external_deterioration"],
            shocked_ca=combined_row["mechanically_shocked_current_account"],
        )

    return quarterly_data, remittance_data, energy_data, qa_baseline_refs, qa_locked_refs

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
GREEN = Font(name=FONT, color="008000")
TITLE = Font(name=FONT, bold=True, size=14)
SECTION = Font(name=FONT, bold=True, size=12, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")
GREY_FILL = PatternFill("solid", fgColor="D9D9D9")
ORANGE_FILL = PatternFill("solid", fgColor="FCE4D6")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
NOTE = Font(name=FONT, italic=True, size=9, color="595959")
RED_BOLD = Font(name=FONT, bold=True, color="C00000")
NUM_FMT = '#,##0.00;(#,##0.00)'
PCT_FMT = '0.0%'
WRAP = Alignment(wrap_text=True, vertical="top")


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================================================
# Sheet 1: README
# ============================================================
def build_readme(wb):
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    set_col_widths(ws, [95])
    r = 1
    ws.cell(r, 1, "External Stress Test — Bangladesh, Pakistan, Vietnam").font = TITLE; r += 2
    ws.cell(r, 1, "Research question").font = BOLD; r += 1
    ws.cell(r, 1, ("How did Bangladesh's exchange-rate and reserve pressure dynamics evolve around the May 2025 move "
        "toward a more flexible exchange-rate arrangement, and what does a simple external-shock stress test suggest "
        "about remaining vulnerabilities? This workbook implements the four-quarter stylized stress-test component.")).alignment = WRAP
    ws.row_dimensions[r].height = 45; r += 2
    ws.cell(r, 1, "This is a stylized mechanical scenario exercise, NOT a forecast.").font = RED_BOLD; r += 1
    ws.cell(r, 1, ("No causal inference, econometric estimation, or forecasting is performed anywhere in this "
        "workbook. Every figure is the arithmetic result of applying a fixed, stated percentage shock to an "
        "observed historical baseline.")).alignment = WRAP
    ws.row_dimensions[r].height = 30; r += 2
    ws.cell(r, 1, "Baseline period and horizon").font = BOLD; r += 1
    ws.cell(r, 1, "Baseline: 2025-Q2 to 2026-Q1 (the latest available common observed quarterly window). Horizon: 4 quarters."); r += 2
    ws.cell(r, 1, "Locked shocks (see Scenarios sheet for the editable assumption cells)").font = BOLD; r += 1
    for s in ["Remittances: -10% (Bangladesh and Pakistan only)",
              "Goods exports: -5% (all three countries)",
              "Global energy-price proxy: +10%, mechanically translated into a stylized +10% increase in goods-import value"]:
        ws.cell(r, 1, "• " + s); r += 1
    r += 1
    ws.cell(r, 1, "On the energy-price proxy").font = BOLD; r += 1
    ws.cell(r, 1, ("The IMF PCPS Energy Index used here is a GLOBAL energy-price shock proxy, not a "
        "country-specific import-price index. Applied identically to all three countries as a stated "
        "simplification.")).alignment = WRAP
    ws.row_dimensions[r].height = 30; r += 2
    ws.cell(r, 1, "Import-price transmission assumption").font = BOLD; r += 1
    ws.cell(r, 1, ('"+10% global energy-price shock -> stylized +10% increase in goods-import value, holding real '
        'import quantities constant." IMPOSED SIMPLIFYING ASSUMPTION, NOT an estimated pass-through coefficient. '
        'The Scenarios sheet shows the energy shock and the transmission rate as separate, auditable cells.')).alignment = WRAP
    ws.row_dimensions[r].height = 45; r += 2
    ws.cell(r, 1, "On Vietnam and remittances").font = BOLD; r += 1
    ws.cell(r, 1, ("Vietnam does NOT receive a remittance shock. No suitable nationwide monthly primary "
        "remittance series was located (UNRESOLVED). Shown as \"N/A\" throughout, never zero, never fabricated.")).alignment = WRAP
    ws.row_dimensions[r].height = 30; r += 2
    ws.cell(r, 1, "On reserves").font = BOLD; r += 1
    ws.cell(r, 1, ("No reserve-loss forecast is produced. Results show mechanically implied cumulative external-"
        "balance deterioration — not a reserve number. Financing, exchange-rate adjustment, capital flows, "
        "valuation effects, and policy response are not modeled.")).alignment = WRAP
    ws.row_dimensions[r].height = 45; r += 2
    ws.cell(r, 1, "Sheets in this workbook").font = BOLD; r += 1
    for s in ["Inputs — observed baseline data, with formula-based totals",
              "Scenarios — editable, clearly labeled assumption cells",
              "Calculations — formula-driven, fully auditable",
              "Summary — reviewer-facing results tables",
              "QA_Checks — automated PASS/FAIL verification"]:
        ws.cell(r, 1, "• " + s); r += 1
    r += 1
    ws.cell(r, 1, "Source / provenance").font = BOLD; r += 1
    for s in ["data/derived/master_analytical_dataset.csv", "data/derived/gate3_stress_test.csv",
              "STRESS_TEST.md", "FINAL_RESULTS_LOCKED.md", "RESEARCH_LEDGER.md"]:
        ws.cell(r, 1, "• " + s); r += 1
    for row in ws.iter_rows(min_row=1, max_row=r, min_col=1, max_col=1):
        for c in row:
            if c.font is None or c.font.name is None:
                c.font = Font(name=FONT)


# ============================================================
# Sheet 2: Inputs — returns key row numbers for downstream sheets
# ============================================================
def build_inputs(wb, quarterly_data, remittance_data, energy_data):
    ws = wb.create_sheet("Inputs")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    set_col_widths(ws, [28, 18, 18, 18])
    ws.cell(1, 1, "INPUTS — Observed Baseline Data (verified, from master_analytical_dataset.csv)").font = TITLE
    ws.cell(2, 1, "Blue = raw observed input. Values USD million unless noted. TOTAL rows use SUM formulas.").font = NOTE

    def section(row, text):
        c = ws.cell(row, 1, text); c.font = SECTION; c.fill = SECTION_FILL
        for col in range(2, 5):
            ws.cell(row, col).fill = SECTION_FILL
        return row + 1

    def qtable_header(row):
        for j, h in enumerate(["Quarter", "Goods Exports", "Goods Imports", "Current Account"], start=1):
            c = ws.cell(row, j, h); c.font = HEADER_FONT; c.fill = PatternFill("solid", fgColor="D9E1F2")
        return row + 1

    quarters = BASELINE_Q
    data = quarterly_data          # read from master_analytical_dataset.csv, not retyped
    rem = remittance_data          # read from master_analytical_dataset.csv, not retyped
    energy = energy_data           # read from master_analytical_dataset.csv, not retyped
    names = {"BGD": "BANGLADESH", "PAK": "PAKISTAN", "VNM": "VIETNAM"}
    row_map = {}

    r = 4
    for iso in ["BGD", "PAK", "VNM"]:
        r = section(r, names[iso])
        r = qtable_header(r)
        first_data_row = r
        for i, q in enumerate(quarters):
            ws.cell(r, 1, q).font = Font(name=FONT)
            for col, key in [(2, "exports"), (3, "imports"), (4, "ca")]:
                c = ws.cell(r, col, data[iso][key][i]); c.font = BLUE; c.number_format = NUM_FMT
            r += 1
        last_data_row = r - 1
        ws.cell(r, 1, "TOTAL (4Q baseline)").font = BOLD
        for col, letter in [(2, 'B'), (3, 'C'), (4, 'D')]:
            c = ws.cell(r, col, f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
            c.font = BOLD; c.number_format = NUM_FMT; c.fill = TOTAL_FILL
        row_map[f"{iso}_quarterly_total_row"] = r
        r += 2

        if iso in rem:
            ws.cell(r, 1, f"{names[iso].title()} Remittances (monthly, for 12-month baseline)").font = Font(name=FONT, bold=True, italic=True); r += 1
            ws.cell(r, 1, "Month").font = HEADER_FONT; ws.cell(r, 1).fill = PatternFill("solid", fgColor="D9E1F2")
            ws.cell(r, 2, "Remittances").font = HEADER_FONT; ws.cell(r, 2).fill = PatternFill("solid", fgColor="D9E1F2")
            r += 1
            first_rem_row = r
            for m, v in rem[iso]:
                ws.cell(r, 1, m).font = Font(name=FONT)
                c = ws.cell(r, 2, v); c.font = BLUE; c.number_format = NUM_FMT
                r += 1
            last_rem_row = r - 1
            ws.cell(r, 1, "TOTAL (12-month baseline)").font = BOLD
            c = ws.cell(r, 2, f"=SUM(B{first_rem_row}:B{last_rem_row})")
            c.font = BOLD; c.number_format = NUM_FMT; c.fill = TOTAL_FILL
            row_map[f"{iso}_remittance_total_row"] = r
            r += 2
        else:
            ws.cell(r, 1, "Vietnam Remittances").font = Font(name=FONT, bold=True, italic=True); r += 1
            c = ws.cell(r, 1, ("N/A — UNRESOLVED. No suitable nationwide monthly primary remittance series was "
                "located for Vietnam (see RESEARCH_LEDGER.md). Not zero. Not fabricated or substituted from "
                "regional (Ho Chi Minh City) or annual data."))
            c.font = RED_BOLD; c.alignment = WRAP
            for col in range(1, 5):
                ws.cell(r, col).fill = ORANGE_FILL
            ws.row_dimensions[r].height = 30
            row_map["VNM_remittance_total_row"] = None
            r += 2

    r = section(r, "GLOBAL ENERGY-PRICE INDEX (contextual input — NOT a country-specific import-price index)")
    ws.cell(r, 1, "Month").font = HEADER_FONT; ws.cell(r, 1).fill = PatternFill("solid", fgColor="D9E1F2")
    ws.cell(r, 2, "Energy Price Index (2016=100)").font = HEADER_FONT; ws.cell(r, 2).fill = PatternFill("solid", fgColor="D9E1F2")
    r += 1
    for m, v in energy:
        ws.cell(r, 1, m).font = Font(name=FONT)
        c = ws.cell(r, 2, v); c.font = BLUE; c.number_format = NUM_FMT
        r += 1
    ws.cell(r, 1, "Source: IMF PCPS, series G001.PNRG.INDEX.M — context only.").font = NOTE

    return row_map  # {"BGD_quarterly_total_row": 10, "BGD_remittance_total_row": 26, "PAK_quarterly_total_row": 34, ...}


# ============================================================
# Sheet 3: Scenarios — returns the row of each assumption cell
# ============================================================
def build_scenarios(wb):
    ws = wb.create_sheet("Scenarios")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    set_col_widths(ws, [48, 14, 55])
    ws.cell(1, 1, "SCENARIOS — Locked Assumptions").font = TITLE
    ws.cell(2, 1, "Yellow cells are the editable SHOCK assumption inputs. All downstream calculations reference "
                  "these cells. The horizon (grey) is a fixed methodological parameter, not an editable scenario "
                  "assumption — the Inputs sheet's baseline is built from a specific, hard-coded set of four "
                  "quarters/twelve months, so changing this cell alone would not actually change which observations "
                  "are pulled.").font = NOTE
    ws.row_dimensions[2].height = 30
    for j, h in enumerate(["Assumption", "Value", "Note"], start=1):
        c = ws.cell(4, j, h); c.font = HEADER_FONT; c.fill = HEADER_FILL

    editable_rows = [
        ("Remittance shock (%)", -0.10, "Applied to Bangladesh and Pakistan only. Vietnam: N/A."),
        ("Export shock (%)", -0.05, "Applied to all three countries."),
        ("Global energy-price shock (%)", 0.10, "Magnitude of the global PCPS energy-price proxy shock. NOT country-specific."),
        ("Import-price transmission / pass-through rate (%)", 1.00, "Real import quantities held constant; 100% pass-through into import VALUE. IMPOSED SIMPLIFYING ASSUMPTION, not an estimated coefficient."),
    ]
    fixed_rows = [
        ("Horizon (quarters)", 4, "FIXED methodological parameter, not an editable scenario assumption — the "
            "baseline is the latest observed 4-quarter/12-month window, a structural choice tied to data "
            "availability, not a shock magnitude. Shown for reference only."),
    ]
    r = 5
    assumption_rows = {}
    for label, val, note in editable_rows:
        ws.cell(r, 1, label).font = Font(name=FONT)
        c = ws.cell(r, 2, val); c.font = Font(name=FONT, bold=True, color="0000FF"); c.fill = YELLOW_FILL
        if isinstance(val, float):
            c.number_format = PCT_FMT
        ws.cell(r, 3, note).font = NOTE
        ws.row_dimensions[r].height = 30
        assumption_rows[label] = r
        r += 1
    for label, val, note in fixed_rows:
        ws.cell(r, 1, label).font = Font(name=FONT)
        c = ws.cell(r, 2, val); c.font = Font(name=FONT, bold=True, color="000000"); c.fill = GREY_FILL
        ws.cell(r, 3, note).font = NOTE
        ws.row_dimensions[r].height = 30
        assumption_rows[label] = r
        r += 1

    r += 1
    ws.cell(r, 1, "Derived (formula, not an additional assumption):").font = BOLD; r += 1
    ws.cell(r, 1, "Import-value shock (%) = Energy shock x Transmission rate")
    energy_row = assumption_rows["Global energy-price shock (%)"]
    transmission_row = assumption_rows["Import-price transmission / pass-through rate (%)"]
    c = ws.cell(r, 2, f"=B{energy_row}*B{transmission_row}"); c.font = Font(name=FONT); c.number_format = PCT_FMT
    import_value_shock_row = r
    r += 2

    ws.cell(r, 1, '"+10% global energy-price shock -> stylized +10% increase in goods-import value, holding real import quantities constant."').font = RED_BOLD
    r += 1
    ws.cell(r, 1, "This is an imposed simplifying assumption, NOT an estimated pass-through coefficient.").font = RED_BOLD

    return {
        "remittance_shock_row": assumption_rows["Remittance shock (%)"],
        "export_shock_row": assumption_rows["Export shock (%)"],
        "import_value_shock_row": import_value_shock_row,
        "horizon_row": assumption_rows["Horizon (quarters)"],
    }


# ============================================================
# Sheet 4: Calculations
# ============================================================
def build_calculations(wb, inputs_rows, scenario_rows):
    ws = wb.create_sheet("Calculations")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B4"
    set_col_widths(ws, [42, 20, 20, 26])
    ws.cell(1, 1, "CALCULATIONS — Fully Formula-Driven (references Inputs and Scenarios sheets)").font = TITLE
    ws.cell(2, 1, "Green = link to another sheet. Recalculates automatically if Inputs or Scenarios change.").font = NOTE
    for j, h in enumerate(["Metric", "Bangladesh", "Pakistan", "Vietnam"], start=1):
        c = ws.cell(4, j, h); c.font = HEADER_FONT; c.fill = HEADER_FILL

    def row_write(r, label, formulas, fmt=NUM_FMT, bold=False):
        ws.cell(r, 1, label).font = BOLD if bold else BLACK
        for col, formula in zip([2, 3, 4], formulas):
            c = ws.cell(r, col, formula)
            c.font = GREEN if (isinstance(formula, str) and formula.startswith("=") and "!" in formula) else (BOLD if bold else BLACK)
            c.number_format = fmt

    r = 6
    ws.cell(r, 1, "Baseline (4-quarter / 12-month totals, linked from Inputs)").font = Font(name=FONT, bold=True, italic=True); r += 1
    row_write(r, "Baseline goods exports",
        [f"=Inputs!B{inputs_rows['BGD_quarterly_total_row']}", f"=Inputs!B{inputs_rows['PAK_quarterly_total_row']}", f"=Inputs!B{inputs_rows['VNM_quarterly_total_row']}"])
    BASELINE_EXP_ROW = r; r += 1
    row_write(r, "Baseline goods imports",
        [f"=Inputs!C{inputs_rows['BGD_quarterly_total_row']}", f"=Inputs!C{inputs_rows['PAK_quarterly_total_row']}", f"=Inputs!C{inputs_rows['VNM_quarterly_total_row']}"])
    BASELINE_IMP_ROW = r; r += 1
    row_write(r, "Baseline current account",
        [f"=Inputs!D{inputs_rows['BGD_quarterly_total_row']}", f"=Inputs!D{inputs_rows['PAK_quarterly_total_row']}", f"=Inputs!D{inputs_rows['VNM_quarterly_total_row']}"])
    BASELINE_CA_ROW = r; r += 1

    ws.cell(r, 1, "Baseline remittances (12mo)").font = BLACK
    c = ws.cell(r, 2, f"=Inputs!B{inputs_rows['BGD_remittance_total_row']}"); c.font = GREEN; c.number_format = NUM_FMT
    c = ws.cell(r, 3, f"=Inputs!B{inputs_rows['PAK_remittance_total_row']}"); c.font = GREEN; c.number_format = NUM_FMT
    c = ws.cell(r, 4, "N/A - UNRESOLVED"); c.font = RED_BOLD
    BASELINE_REM_ROW = r; r += 2

    ws.cell(r, 1, "Scenario shock magnitudes (linked from Scenarios)").font = Font(name=FONT, bold=True, italic=True); r += 1
    # Remittance shock % — Vietnam shown as "N/A", not the -10% assumption, since Vietnam does not
    # receive this shock in the locked Gate 3 methodology (no remittance baseline exists for it).
    ws.cell(r, 1, "Remittance shock (%)").font = BLACK
    c = ws.cell(r, 2, f"=Scenarios!$B${scenario_rows['remittance_shock_row']}"); c.font = GREEN; c.number_format = PCT_FMT
    c = ws.cell(r, 3, f"=Scenarios!$B${scenario_rows['remittance_shock_row']}"); c.font = GREEN; c.number_format = PCT_FMT
    c = ws.cell(r, 4, "N/A"); c.font = RED_BOLD
    REM_PCT_ROW = r; r += 1
    row_write(r, "Export shock (%)", [f"=Scenarios!$B${scenario_rows['export_shock_row']}"] * 3, fmt=PCT_FMT)
    EXP_PCT_ROW = r; r += 1
    row_write(r, "Import-value shock (%) [energy x transmission]", [f"=Scenarios!$B${scenario_rows['import_value_shock_row']}"] * 3, fmt=PCT_FMT)
    IMP_PCT_ROW = r; r += 2

    ws.cell(r, 1, "Shock amounts ($ million; negative = deterioration)").font = Font(name=FONT, bold=True, italic=True); r += 1
    ws.cell(r, 1, "Remittance shock amount").font = BLACK
    for col, letter in [(2, 'B'), (3, 'C'), (4, 'D')]:
        formula = f'=IF(ISNUMBER({letter}{BASELINE_REM_ROW}),{letter}{BASELINE_REM_ROW}*{letter}{REM_PCT_ROW},"N/A")'
        c = ws.cell(r, col, formula)
        c.font = RED_BOLD if letter == 'D' else BLACK
        c.number_format = NUM_FMT
    REM_SHOCK_ROW = r; r += 1

    row_write(r, "Export shock amount",
        [f"=B{BASELINE_EXP_ROW}*B{EXP_PCT_ROW}", f"=C{BASELINE_EXP_ROW}*C{EXP_PCT_ROW}", f"=D{BASELINE_EXP_ROW}*D{EXP_PCT_ROW}"])
    EXP_SHOCK_ROW = r; r += 1
    row_write(r, "Import-price shock amount",
        [f"=-1*B{BASELINE_IMP_ROW}*B{IMP_PCT_ROW}", f"=-1*C{BASELINE_IMP_ROW}*C{IMP_PCT_ROW}", f"=-1*D{BASELINE_IMP_ROW}*D{IMP_PCT_ROW}"])
    IMP_SHOCK_ROW = r; r += 2

    ws.cell(r, 1, "Combined deterioration and mechanically shocked current account").font = Font(name=FONT, bold=True, italic=True); r += 1
    ws.cell(r, 1, "Combined external-account deterioration").font = BOLD
    c = ws.cell(r, 2, f"=B{REM_SHOCK_ROW}+B{EXP_SHOCK_ROW}+B{IMP_SHOCK_ROW}"); c.font = BOLD; c.number_format = NUM_FMT; c.fill = TOTAL_FILL
    c = ws.cell(r, 3, f"=C{REM_SHOCK_ROW}+C{EXP_SHOCK_ROW}+C{IMP_SHOCK_ROW}"); c.font = BOLD; c.number_format = NUM_FMT; c.fill = TOTAL_FILL
    c = ws.cell(r, 4, f"=D{EXP_SHOCK_ROW}+D{IMP_SHOCK_ROW}"); c.font = BOLD; c.number_format = NUM_FMT; c.fill = TOTAL_FILL  # Vietnam: no remittance term
    COMBINED_ROW = r; r += 1

    row_write(r, "Mechanically shocked current account",
        [f"=B{BASELINE_CA_ROW}+B{COMBINED_ROW}", f"=C{BASELINE_CA_ROW}+C{COMBINED_ROW}", f"=D{BASELINE_CA_ROW}+D{COMBINED_ROW}"], bold=True)
    SHOCKED_CA_ROW = r; r += 1

    row_write(r, "Horizon (quarters)", [f"=Scenarios!$B${scenario_rows['horizon_row']}"] * 3, fmt='0')
    HORIZON_ROW = r; r += 1

    return dict(BASELINE_EXP_ROW=BASELINE_EXP_ROW, BASELINE_IMP_ROW=BASELINE_IMP_ROW, BASELINE_CA_ROW=BASELINE_CA_ROW,
                BASELINE_REM_ROW=BASELINE_REM_ROW, REM_PCT_ROW=REM_PCT_ROW, REM_SHOCK_ROW=REM_SHOCK_ROW,
                EXP_SHOCK_ROW=EXP_SHOCK_ROW, IMP_SHOCK_ROW=IMP_SHOCK_ROW, COMBINED_ROW=COMBINED_ROW,
                SHOCKED_CA_ROW=SHOCKED_CA_ROW, HORIZON_ROW=HORIZON_ROW)


# ============================================================
# Sheet 5: Summary
# ============================================================
def build_summary(wb, calc_rows):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    # Column B and C widened so the lower table's scenario labels (e.g. "Import-price shock only") and its
    # long "Mechanically shocked current account" header are fully visible rather than truncated.
    set_col_widths(ws, [16, 28, 40, 16, 18, 24, 26])
    ws.cell(1, 1, "SUMMARY — Reviewer-Facing Results (all figures link to Calculations sheet)").font = TITLE

    r = 3
    headers = ["Country", "Baseline CA", "Remittance shock", "Export shock", "Import-price shock", "Combined deterioration", "Mechanically shocked CA"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(r, j, h); c.font = HEADER_FONT; c.fill = HEADER_FILL
    r += 1
    for name, col in [("Bangladesh", "B"), ("Pakistan", "C"), ("Vietnam", "D")]:
        ws.cell(r, 1, name).font = Font(name=FONT, bold=True)
        ws.cell(r, 2, f"=Calculations!{col}{calc_rows['BASELINE_CA_ROW']}").font = GREEN
        ws.cell(r, 3, f"=Calculations!{col}{calc_rows['REM_SHOCK_ROW']}").font = GREEN
        ws.cell(r, 4, f"=Calculations!{col}{calc_rows['EXP_SHOCK_ROW']}").font = GREEN
        ws.cell(r, 5, f"=Calculations!{col}{calc_rows['IMP_SHOCK_ROW']}").font = GREEN
        ws.cell(r, 6, f"=Calculations!{col}{calc_rows['COMBINED_ROW']}").font = BOLD; ws.cell(r, 6).fill = TOTAL_FILL
        ws.cell(r, 7, f"=Calculations!{col}{calc_rows['SHOCKED_CA_ROW']}").font = BOLD; ws.cell(r, 7).fill = TOTAL_FILL
        for cc in range(2, 8):
            ws.cell(r, cc).number_format = NUM_FMT
        r += 1

    ws.cell(r, 1, "Vietnam combined scenario = Export + import-price only").font = RED_BOLD; r += 1
    ws.cell(r, 1, "Note: Vietnam does not include a remittance shock because no suitable nationwide monthly remittance series was available.").font = RED_BOLD
    ws.row_dimensions[r].height = 18; r += 3

    ws.cell(r, 1, "Bangladesh and Pakistan — individual scenario detail (mechanically shocked CA under each scenario)").font = Font(name=FONT, bold=True, size=11); r += 1
    for j, h in enumerate(["Country", "Scenario", "Mechanically shocked current account"], start=1):
        c = ws.cell(r, j, h); c.font = HEADER_FONT; c.fill = HEADER_FILL
    r += 1
    scen_defs = [("Baseline", None), ("Remittance shock only", calc_rows['REM_SHOCK_ROW']),
                 ("Export shock only", calc_rows['EXP_SHOCK_ROW']), ("Import-price shock only", calc_rows['IMP_SHOCK_ROW']),
                 ("Combined", calc_rows['COMBINED_ROW'])]
    for name, col in [("Bangladesh", "B"), ("Pakistan", "C")]:
        for scen_label, shock_row in scen_defs:
            ws.cell(r, 1, name).font = Font(name=FONT)
            ws.cell(r, 2, scen_label).font = Font(name=FONT)
            formula = f"=Calculations!{col}{calc_rows['BASELINE_CA_ROW']}" if shock_row is None else f"=Calculations!{col}{calc_rows['BASELINE_CA_ROW']}+Calculations!{col}{shock_row}"
            c = ws.cell(r, 3, formula); c.font = GREEN; c.number_format = NUM_FMT
            if scen_label == "Combined":
                c.font = Font(name=FONT, bold=True, color="008000"); ws.cell(r, 3).fill = TOTAL_FILL
            r += 1
        r += 1


# ============================================================
# Sheet 6: QA_Checks
# ============================================================
def build_qa_checks(wb, inputs_rows, calc_rows, scenario_rows, qa_baseline_refs, qa_locked_refs):
    ws = wb.create_sheet("QA_Checks")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    set_col_widths(ws, [55, 20, 20, 45])
    ws.cell(1, 1, "QA_CHECKS — Automated Verification").font = TITLE
    ws.cell(2, 1, "All checks are formulas. Reference values (blue) are cited from the project's finalized files.").font = NOTE
    for j, h in enumerate(["Check", "Result", "Reference value", "Note"], start=1):
        c = ws.cell(4, j, h); c.font = HEADER_FONT; c.fill = HEADER_FILL

    r = 5
    check_rows = []

    def add_check(label, formula, note=""):
        nonlocal r
        ws.cell(r, 1, label).font = Font(name=FONT)
        ws.cell(r, 2, formula).font = Font(name=FONT, bold=True)
        ws.cell(r, 4, note).font = NOTE
        check_rows.append(r)
        r += 1

    b, i = inputs_rows, calc_rows

    # Quarterly/remittance data ranges derived from inputs_rows (the actual sheet-building row
    # positions), rather than re-hardcoded literal row numbers duplicating what build_inputs already
    # knows. Quarterly blocks are always 4 rows immediately above their TOTAL row; remittance blocks
    # are always 12 rows immediately above theirs.
    def quarterly_range(total_row):
        return total_row - 4, total_row - 1

    def remittance_range(total_row):
        return total_row - 12, total_row - 1

    for name, total_key in [("Bangladesh", "BGD_quarterly_total_row"), ("Pakistan", "PAK_quarterly_total_row"), ("Vietnam", "VNM_quarterly_total_row")]:
        first, last = quarterly_range(b[total_key])
        add_check(f"1. {name}: 4 quarterly observations exist",
            f'=IF(AND(COUNT(Inputs!B{first}:B{last})=4,COUNT(Inputs!C{first}:C{last})=4,COUNT(Inputs!D{first}:D{last})=4),"PASS","FAIL")')

    for name, total_key in [("Bangladesh", "BGD_remittance_total_row"), ("Pakistan", "PAK_remittance_total_row")]:
        first, last = remittance_range(b[total_key])
        add_check(f"2. {name}: 12 monthly remittance observations exist", f'=IF(COUNT(Inputs!B{first}:B{last})=12,"PASS","FAIL")')
    add_check("3. Vietnam remittances are marked N/A, not zero",
        f'=IF(AND(ISTEXT(Calculations!D{i["BASELINE_REM_ROW"]}),ISTEXT(Calculations!D{i["REM_SHOCK_ROW"]}),NOT(ISNUMBER(Calculations!D{i["REM_SHOCK_ROW"]}))),"PASS","FAIL")')
    add_check("3b. Vietnam remittance shock assumption (%) displays as text \"N/A\", not a numeric percentage",
        f'=IF(AND(ISTEXT(Calculations!D{i["REM_PCT_ROW"]}),Calculations!D{i["REM_PCT_ROW"]}="N/A",NOT(ISNUMBER(Calculations!D{i["REM_PCT_ROW"]}))),"PASS","FAIL")')
    add_check("4. Bangladesh: combined = sum of individual shocks",
        f'=IF(ABS(Calculations!B{i["COMBINED_ROW"]}-(Calculations!B{i["REM_SHOCK_ROW"]}+Calculations!B{i["EXP_SHOCK_ROW"]}+Calculations!B{i["IMP_SHOCK_ROW"]}))<0.01,"PASS","FAIL")')
    add_check("4. Pakistan: combined = sum of individual shocks",
        f'=IF(ABS(Calculations!C{i["COMBINED_ROW"]}-(Calculations!C{i["REM_SHOCK_ROW"]}+Calculations!C{i["EXP_SHOCK_ROW"]}+Calculations!C{i["IMP_SHOCK_ROW"]}))<0.01,"PASS","FAIL")')
    add_check("4. Vietnam: combined = sum of applicable shocks (export+import-price only)",
        f'=IF(ABS(Calculations!D{i["COMBINED_ROW"]}-(Calculations!D{i["EXP_SHOCK_ROW"]}+Calculations!D{i["IMP_SHOCK_ROW"]}))<0.01,"PASS","FAIL")')
    for name, col in [("Bangladesh", "B"), ("Pakistan", "C"), ("Vietnam", "D")]:
        add_check(f"5. {name}: shocked CA = baseline CA + combined deterioration",
            f'=IF(ABS(Calculations!{col}{i["SHOCKED_CA_ROW"]}-(Calculations!{col}{i["BASELINE_CA_ROW"]}+Calculations!{col}{i["COMBINED_ROW"]}))<0.01,"PASS","FAIL")')
    add_check("6. Horizon equals 4 quarters (all countries)",
        f'=IF(AND(Calculations!B{i["HORIZON_ROW"]}=4,Calculations!C{i["HORIZON_ROW"]}=4,Calculations!D{i["HORIZON_ROW"]}=4,Scenarios!B{scenario_rows["horizon_row"]}=4),"PASS","FAIL")')

    ws.cell(r, 1, "7. Baseline inputs match data/derived/gate3_stress_test.csv").font = Font(name=FONT, bold=True, italic=True); r += 1
    # Reference values read directly from gate3_stress_test.csv (qa_baseline_refs), not retyped constants.
    ref_vals = []
    for name, col in [("Bangladesh", "B"), ("Pakistan", "C"), ("Vietnam", "D")]:
        iso = {"B": "BGD", "C": "PAK", "D": "VNM"}[col]
        b = qa_baseline_refs[iso]
        ref_vals.append((f"{name} baseline exports", f"{col}{i['BASELINE_EXP_ROW']}", b["exports"]))
        ref_vals.append((f"{name} baseline imports", f"{col}{i['BASELINE_IMP_ROW']}", b["imports"]))
        ref_vals.append((f"{name} baseline CA", f"{col}{i['BASELINE_CA_ROW']}", b["ca"]))
        if b["remittances"] is not None:
            ref_vals.append((f"{name} baseline remittances", f"{col}{i['BASELINE_REM_ROW']}", b["remittances"]))
    for label, calc_cell, ref_val in ref_vals:
        ws.cell(r, 1, f"    {label}").font = Font(name=FONT)
        c = ws.cell(r, 3, ref_val); c.font = BLUE; c.number_format = NUM_FMT
        ws.cell(r, 2, f'=IF(ROUND(Calculations!{calc_cell},2)=ROUND(C{r},2),"PASS","FAIL")').font = Font(name=FONT, bold=True)
        ws.cell(r, 4, "Source: gate3_stress_test.csv").font = NOTE
        check_rows.append(r); r += 1

    ws.cell(r, 1, "8. Workbook results match the verified Gate 3 reference values (from gate3_stress_test.csv, "
                  "consistent with FINAL_RESULTS_LOCKED.md R-005)").font = Font(name=FONT, bold=True, italic=True); r += 1
    # Reference values read directly from gate3_stress_test.csv (qa_locked_refs), not retyped constants.
    locked_vals = []
    for name, col in [("Bangladesh", "B"), ("Pakistan", "C"), ("Vietnam", "D")]:
        iso = {"B": "BGD", "C": "PAK", "D": "VNM"}[col]
        loc = qa_locked_refs[iso]
        suffix = " (2-shock)" if iso == "VNM" else ""
        locked_vals.append((f"{name} combined deterioration{suffix}", f"{col}{i['COMBINED_ROW']}", loc["deterioration"]))
        locked_vals.append((f"{name} shocked CA", f"{col}{i['SHOCKED_CA_ROW']}", loc["shocked_ca"]))
    for label, calc_cell, ref_val in locked_vals:
        ws.cell(r, 1, f"    {label}").font = Font(name=FONT)
        c = ws.cell(r, 3, ref_val); c.font = BLUE; c.number_format = NUM_FMT
        ws.cell(r, 2, f'=IF(ROUND(Calculations!{calc_cell},2)=ROUND(C{r},2),"PASS","FAIL")').font = Font(name=FONT, bold=True)
        ws.cell(r, 4, "Source: gate3_stress_test.csv (consistent with FINAL_RESULTS_LOCKED.md R-005)").font = NOTE
        check_rows.append(r); r += 1

    r += 1
    ws.cell(r, 1, "OVERALL RESULT").font = Font(name=FONT, bold=True, size=12)
    first_chk, last_chk = min(check_rows), max(check_rows)
    c = ws.cell(r, 2, f'=IF(COUNTIF(B{first_chk}:B{last_chk},"FAIL")=0,"ALL CHECKS PASSED","REVIEW REQUIRED")')
    c.font = Font(name=FONT, bold=True, size=12); c.fill = PatternFill("solid", fgColor="C6EFCE")

    # Conditional formatting
    green_fill = PatternFill("solid", fgColor="C6EFCE"); green_font = Font(name=FONT, bold=True, color="006100")
    red_fill = PatternFill("solid", fgColor="FFC7CE"); red_font = Font(name=FONT, bold=True, color="9C0006")
    for value in ["PASS", "ALL CHECKS PASSED"]:
        ws.conditional_formatting.add(f"B5:B{r}", CellIsRule(operator="equal", formula=[f'"{value}"'], fill=green_fill, font=green_font))
    for value in ["FAIL", "REVIEW REQUIRED"]:
        ws.conditional_formatting.add(f"B5:B{r}", CellIsRule(operator="equal", formula=[f'"{value}"'], fill=red_fill, font=red_font))


def main():
    quarterly_data, remittance_data, energy_data, qa_baseline_refs, qa_locked_refs = load_workbook_data()

    wb = openpyxl.Workbook()
    build_readme(wb)
    inputs_rows = build_inputs(wb, quarterly_data, remittance_data, energy_data)
    scenario_rows = build_scenarios(wb)
    calc_rows = build_calculations(wb, inputs_rows, scenario_rows)
    build_summary(wb, calc_rows)
    build_qa_checks(wb, inputs_rows, calc_rows, scenario_rows, qa_baseline_refs, qa_locked_refs)

    tab_colors = {"README": "1F4E79", "Inputs": "2E7D32", "Scenarios": "BF8F00",
                  "Calculations": "1F4E79", "Summary": "C00000", "QA_Checks": "595959"}
    for name, color in tab_colors.items():
        wb[name].sheet_properties.tabColor = color

    import os
    os.makedirs("stress_test", exist_ok=True)
    wb.save("stress_test/External_Stress_Test.xlsx")
    print("Workbook built: stress_test/External_Stress_Test.xlsx")
    print("Next: recalculate (e.g. with the xlsx skill's recalc.py) and confirm QA_Checks reads ALL CHECKS PASSED.")


if __name__ == "__main__":
    main()
